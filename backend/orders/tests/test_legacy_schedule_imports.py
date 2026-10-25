"""Legacy supplier exports use demand identities, not the repeated sales IDs."""

import tempfile
from datetime import date
from pathlib import Path

from django.contrib.auth import get_user_model
from django.test import override_settings
from openpyxl import Workbook
from rest_framework.test import APITestCase

from orders.imports import commit_business_batch, preview_business_workbook
from orders.models import BusinessImportBatch, ProductSpecification
from quality.models import QualityOrder

from .test_imports import factory_workbook, upload, workbook_bytes


def legacy_schedule(*, quantity=336010, repeat_header=False):
    workbook = Workbook()
    headers = [
        "订单号", "订单项次", "需求单号", "项次", "规格", "B胶配方名",
        "工厂交期", "生产预交日期", "订单量", "应生产数量", "已下单量",
        "订单日期", "成型工时", "需求胶料重量", "模具编号", "结案否",
    ]
    for index, name in enumerate(("NBR-T3", "NBR-T5")):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = name
        sheet.append(headers)
        sheet.append([
            "SALES-OLD", 2, "DEMAND-OLD", "03", "TEST-LEGACY-A", "TEST-MATERIAL",
            date(2026, 4, 24), date(2026, 4, 20), 500000, 336010,
            quantity, date(2026, 3, 28), 116, 565, "MOLD-OLD", "N",
        ])
        if index == 0:
            if repeat_header:
                sheet.append(headers)
            sheet.append([
                "SALES-OLD", 2, "DEMAND-OLD", "04", "TEST-LEGACY-A", "TEST-MATERIAL",
                date(2026, 4, 24), date(2026, 4, 20), 500000, 190306,
                190306, date(2026, 3, 28), 66, 320, "MOLD-OLD", "N",
            ])
        else:
            sheet["C2"] = "DEMAND-OTHER"
            sheet["E2"] = "TEST-LEGACY-AB"
    return workbook_bytes(workbook)


class LegacyScheduleImportTests(APITestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="legacy-importer", password="test")
        self.client.force_authenticate(self.user)
        self.media = tempfile.TemporaryDirectory()
        self.settings_override = override_settings(MEDIA_ROOT=Path(self.media.name))
        self.settings_override.enable()
        self.addCleanup(self.media.cleanup)
        self.addCleanup(self.settings_override.disable)

    def preview(self, content=None, **filters):
        return preview_business_workbook(upload("3-30.xlsx", content or legacy_schedule()), self.user, **filters)

    def commit(self, preview):
        return commit_business_batch(
            BusinessImportBatch.objects.get(pk=preview["token"]), self.user,
            cut_length_acknowledged=True,
            cut_length_digest=(preview.get("cut_length_check") or {}).get("digest", ""),
        )

    def test_multiple_sheets_keep_demand_items_and_allocated_quantities(self):
        preview = self.preview(legacy_schedule(repeat_header=True))
        self.assertEqual(preview["error_count"], 0, preview["issues"])
        self.assertEqual(preview["counts"]["orders"], 3)
        self.commit(preview)
        orders = QualityOrder.objects.filter(order_no="DEMAND-OLD").order_by("item_no")
        self.assertEqual(list(orders.values_list("item_no", "order_quantity")), [("03", 336010), ("04", 190306)])
        self.assertEqual(orders[0].source_system, "NBR-T3")
        self.assertEqual(orders[0].order_date, date(2026, 3, 28))
        self.assertEqual(orders[0].due_date, date(2026, 4, 24))
        self.assertIsNone(orders[0].source_document_at)
        self.assertEqual(orders[0].source_row, 2)
        self.assertEqual(orders[0].raw_data["订单号"]["raw_value"], "SALES-OLD")
        self.assertEqual(orders[0].raw_data["订单量"]["raw_value"], 500000)
        self.assertEqual(orders[0].production_quantity, "")
        self.assertEqual(orders[0].shipped_quantity, "")
        self.assertEqual(orders[0].product_specification_id, orders[1].product_specification_id)
        self.assertEqual(QualityOrder.objects.get(order_no="DEMAND-OTHER").source_system, "NBR-T5")

    def test_exact_scope_does_not_import_other_products_and_reimport_is_idempotent(self):
        filters = {"specification": " test-legacy-a ", "material": "test-material"}
        preview = self.preview(**filters)
        self.assertEqual(preview["error_count"], 0, preview["issues"])
        self.assertEqual(preview["counts"]["orders"], 2)
        self.assertEqual(preview["counts"]["product_specifications"], 1)
        self.commit(preview)
        second = self.preview(**filters)
        result = self.commit(second)
        self.assertEqual(result["created"]["orders"], 0)
        self.assertEqual(QualityOrder.objects.count(), 2)
        self.assertEqual(ProductSpecification.objects.count(), 1)
        self.assertFalse(QualityOrder.objects.filter(order_no="DEMAND-OTHER").exists())
        batch = BusinessImportBatch.objects.get(pk=preview["token"])
        self.assertEqual(batch.payload["filters"], filters)

    def test_missing_or_zero_demand_quantity_never_falls_back_to_sales_total(self):
        for quantity in (None, 0, -1, "invalid"):
            with self.subTest(quantity=quantity):
                preview = self.preview(legacy_schedule(quantity=quantity))
                self.assertGreater(preview["error_count"], 0)
                with self.assertRaisesMessage(ValueError, "预检存在错误"):
                    self.commit(preview)
        self.assertEqual(QualityOrder.objects.count(), 0)

    def test_legacy_missing_process_columns_preserve_current_product_and_manual_fields(self):
        product = ProductSpecification.objects.create(
            specification="TEST-LEGACY-A", material="TEST-MATERIAL", mold_no="MOLD-OLD",
            material_length="300", cut_weight="21", primary_curing="60s",
            actual_material_length="310", actual_cut_weight="22", mold_size="500x500",
        )
        preview = self.preview(specification="TEST-LEGACY-A")
        self.assertEqual(preview["error_count"], 0, preview["issues"])
        self.commit(preview)
        product.refresh_from_db()
        self.assertEqual((product.material_length, product.cut_weight, product.primary_curing), ("300", "21", "60s"))
        self.assertEqual((product.actual_material_length, product.actual_cut_weight, product.mold_size), ("310", "22", "500x500"))

    def test_scope_keeps_linked_criteria_in_current_factory_format(self):
        preview = self.preview(
            factory_workbook(specification="TEST-LEGACY-A"),
            specification="TEST-LEGACY-A",
        )
        self.assertEqual(preview["error_count"], 0, preview["issues"])
        self.assertEqual(preview["counts"]["inspection_criteria"], 1)
        self.commit(preview)

    def test_no_match_fails_instead_of_importing_whole_workbook(self):
        with self.assertRaisesMessage(ValueError, "未找到符合指定规格"):
            self.preview(specification="TEST-LEGACY")
        self.assertEqual(QualityOrder.objects.count(), 0)

    def test_legacy_header_whitespace_is_normalized_consistently(self):
        from openpyxl import load_workbook
        import io
        workbook = load_workbook(io.BytesIO(legacy_schedule()))
        for sheet in workbook:
            sheet["F1"] = " B胶\n配方名 "
            sheet["C1"] = " 需求单号 "
            sheet["K1"] = "已下单量\n"
        preview = self.preview(workbook_bytes(workbook), specification="TEST-LEGACY-A")
        self.assertEqual(preview["error_count"], 0, preview["issues"])
        self.commit(preview)
        self.assertEqual(QualityOrder.objects.count(), 2)

    def test_scope_keeps_document_date_conflict_warning(self):
        preview = preview_business_workbook(
            upload("factory-2026-08-05.xlsx", factory_workbook(specification="TEST-LEGACY-A")),
            self.user, specification="TEST-LEGACY-A",
        )
        self.assertTrue(any("文件名识别的" in issue["message"] for issue in preview["issues"]), preview["issues"])

    def test_old_schedule_cannot_revert_a_more_recent_customer_document(self):
        from django.utils import timezone
        product = ProductSpecification.objects.create(
            specification="TEST-LEGACY-A", material="TEST-MATERIAL", mold_no="MOLD-OLD",
            primary_curing="NEW-CURING",
        )
        order = QualityOrder.objects.create(
            order_no="DEMAND-OLD", item_no="03", specification="TEST-LEGACY-A",
            material="TEST-MATERIAL", order_quantity=123, source_system="NBR-T3",
            source_document_at=timezone.now(), product_specification=product, created_by=self.user,
        )
        preview = self.preview(specification="TEST-LEGACY-A")
        self.assertEqual(preview["error_count"], 0, preview["issues"])
        self.commit(preview)
        order.refresh_from_db()
        product.refresh_from_db()
        self.assertEqual(order.order_quantity, 123)
        self.assertEqual(product.primary_curing, "NEW-CURING")
        self.assertEqual(QualityOrder.objects.filter(order_no="DEMAND-OLD").count(), 2)

    def test_nonmatching_rows_cannot_block_or_expand_scoped_import(self):
        from openpyxl import load_workbook
        import io
        workbook = load_workbook(io.BytesIO(legacy_schedule()))
        workbook["NBR-T5"]["K2"] = "bad-quantity"
        preview = self.preview(workbook_bytes(workbook), specification="TEST-LEGACY-A")
        self.assertEqual(preview["error_count"], 0, preview["issues"])
        self.commit(preview)
        self.assertEqual(QualityOrder.objects.count(), 2)

    def test_api_applies_scope_before_commit_and_keeps_original_file(self):
        content = legacy_schedule()
        response = self.client.post("/api/orders/imports/preview/", {
            "file": upload("3-30.xlsx", content),
            "specification": "TEST-LEGACY-A", "material": "TEST-MATERIAL",
        }, format="multipart")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["counts"]["orders"], 2)
        batch = BusinessImportBatch.objects.get(pk=response.data["token"])
        with batch.original_file.open("rb") as original:
            self.assertEqual(original.read(), content)
        response = self.client.post("/api/orders/imports/commit/", {
            "token": str(batch.pk), "cut_length_acknowledged": True,
            "cut_length_digest": response.data["cut_length_check"]["digest"],
        }, format="json")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(QualityOrder.objects.count(), 2)
